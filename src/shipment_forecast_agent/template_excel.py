"""Match forecasts to original rows; never reconstruct the tender workbook."""

import math
from copy import copy
from datetime import date

from openpyxl import load_workbook

from .excel_io import DELIVERY_REGION_ALIASES, SHIPPING_POINT_ALIASES

REQUIRED_TENDER_HEADERS = (
    "№ направления",
    "Наименование пункта отгрузки",
    "График работы пункта отгрузки",
    "Код страны назначения",
    "Наименование региона доставки",
    "Код региона доставки",
    "Адрес пункта доставки",
    "Код зоны доставки Логинет",
    "График работы пункта доставки",
    "Валюта",
    "Наименование продуктов",
    "Класс опасности",
    "Вид упаковки",
    "Вид и тип ТС",
    "Кол-во груза загружаемое в ед. ТС, т",
    "Комментарии заказчика",
    "Целевая ставка",
    "БП",
    "Базовая цена",
    "Цена за единицу ТС, без НДС",
    "Гарантированные объемы, единиц ТС\\месяц",
)

FORECAST_COLUMNS = (
    ("forecast_m1", "Прогноз M1"),
    ("forecast_m2", "Прогноз M2"),
    ("forecast_m3", "Прогноз M3"),
)


def normalized(value):
    return " ".join(str(value or "").split()).casefold()


def current_month():
    """Calendar month on the machine running the agent, frozen in request.json."""
    return date.today().strftime("%Y-%m")  # noqa: DTZ011 - local business month


def validate_tender_headers(book, settings):
    """Validate the supplied tender template without requiring optional values."""
    sheet = book[settings.EXCEL_SHEET] if settings.EXCEL_SHEET else book.active
    headers = [normalized(c.value) for c in sheet[settings.EXCEL_HEADER_ROW]]
    missing, duplicates = [], []
    for label in REQUIRED_TENDER_HEADERS:
        aliases = {normalized(label)}
        if label == "Наименование пункта отгрузки":
            aliases |= SHIPPING_POINT_ALIASES | {
                "shipment_point_name",
                "пункт выгрузки",
            }
        elif label == "Наименование региона доставки":
            aliases |= DELIVERY_REGION_ALIASES | {"delivery_point_name"}
        count = sum(h in aliases for h in headers)
        if count == 0:
            missing.append(label)
        elif count > 1:
            duplicates.append(label)
    errors = []
    if missing:
        errors.append("Отсутствуют обязательные столбцы: " + "; ".join(missing))
    if duplicates:
        errors.append("Дублируются обязательные столбцы: " + "; ".join(duplicates))
    if errors:
        raise ValueError(" | ".join(errors))


def validate_tender(path, settings):
    book = load_workbook(path)
    try:
        validate_tender_headers(book, settings)
        rows(book, settings)
    finally:
        book.close()


def rows(book, settings):
    sheet = book[settings.EXCEL_SHEET] if settings.EXCEL_SHEET else book.active
    headers = {
        normalized(c.value): c.column
        for c in sheet[settings.EXCEL_HEADER_ROW]
        if c.value
    }

    def find(aliases):
        matches = [headers[a] for a in aliases if a in headers]
        if len(matches) != 1:
            raise ValueError("Expected exactly one point column and one region column")
        return matches[0]

    point = find(SHIPPING_POINT_ALIASES | {"shipment_point_name", "пункт выгрузки"})
    region = find(DELIVERY_REGION_ALIASES | {"delivery_point_name"})
    result = []
    for index in range(settings.EXCEL_HEADER_ROW + 1, sheet.max_row + 1):
        p, r = sheet.cell(index, point).value, sheet.cell(index, region).value
        if not p and not r:
            continue
        if not p or not r:
            raise ValueError(f"Incomplete direction in row {index}")
        if any(isinstance(v, str) and v.startswith("=") for v in (p, r)):
            raise ValueError("Direction columns must contain values, not formulas")
        result.append((index, str(p).strip(), str(r).strip()))
    if not result:
        raise ValueError("No directions in workbook")
    return sheet, result


def request_data(path, settings, request_id):
    book = load_workbook(path)
    try:
        validate_tender_headers(book, settings)
        _, source = rows(book, settings)
        directions = list(dict.fromkeys((r[1], r[2]) for r in source))
        if len(directions) != 1:
            raise ValueError(
                "Input JSON schema supports exactly one unique direction; "
                f"found {len(directions)}"
            )
        point, region = directions[0]
        return {
            "mail_id": request_id,
            "param_month_m0": current_month(),
            "shipment_point_name": point,
            "delivery_point_name": region,
        }
    finally:
        book.close()


def fill_forecasts(source, output, forecasts, settings):
    book = load_workbook(source)
    try:
        sheet, source_rows = rows(book, settings)
        predictions = {}
        for item in forecasts:
            p = item.get("shipment_point_name", item.get("shipping_point"))
            r = item.get("delivery_point_name", item.get("delivery_region"))
            if not p or not r:
                raise ValueError("Missing forecast direction")
            values = []
            for field, _ in FORECAST_COLUMNS:
                raw_forecast = item.get(field)
                if isinstance(raw_forecast, bool) or raw_forecast is None:
                    raise ValueError(f"Missing or invalid {field}")
                try:
                    value = float(raw_forecast)
                except (TypeError, ValueError):
                    raise ValueError(f"Missing or invalid {field}") from None
                if not math.isfinite(value) or value < 0:
                    raise ValueError(f"{field} must be finite and nonnegative")
                values.append(value)
            key = normalized(p), normalized(r)
            if key in predictions:
                raise ValueError(
                    "Multiple forecasts for one direction; ambiguous response"
                )
            predictions[key] = tuple(values)
        values, used = [], set()
        for index, p, r in source_rows:
            key = normalized(p), normalized(r)
            if key not in predictions:
                raise ValueError(f"Missing forecast for row {index}")
            values.append((index, predictions[key]))
            used.add(key)
        if used != set(predictions):
            raise ValueError("Response contains directions absent from Excel")
        existing_headers = {
            normalized(c.value) for c in sheet[settings.EXCEL_HEADER_ROW] if c.value
        }
        result_headers = {normalized(label) for _, label in FORECAST_COLUMNS}
        if "прогноз" in existing_headers or existing_headers & result_headers:
            raise ValueError(
                "Source already contains forecast columns; refusing overwrite"
            )
        first_column = sheet.max_column + 1
        style_column = first_column - 1
        for offset, (_, header) in enumerate(FORECAST_COLUMNS):
            column = first_column + offset
            column_values = [
                (index, forecast_values[offset]) for index, forecast_values in values
            ]
            for index, value in [
                (settings.EXCEL_HEADER_ROW, header),
                *column_values,
            ]:
                target = sheet.cell(index, column, value)
                target._style = copy(sheet.cell(index, style_column)._style)
                if index != settings.EXCEL_HEADER_ROW:
                    target.number_format = "0.00"
        output.parent.mkdir(parents=True, exist_ok=True)
        book.save(output)
        return output
    finally:
        book.close()
