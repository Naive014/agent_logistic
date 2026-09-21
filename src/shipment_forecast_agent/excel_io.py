import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import Direction, ForecastResponse

SHIPPING_POINT_ALIASES = {
    "пункт отгрузки",
    "наименование пункта отгрузки",
    "shipping_point",
}
DELIVERY_REGION_ALIASES = {
    "регион доставки",
    "наименование региона доставки",
    "delivery_region",
}


class InputWorkbookError(ValueError):
    pass


def _normalize_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _find_column(headers: list[str], aliases: set[str], label: str) -> int:
    for index, header in enumerate(headers):
        if header in aliases:
            return index
    raise InputWorkbookError(f"Missing required column: {label}")


def read_directions(path: str | Path) -> list[Direction]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return _read_sheet(workbook.active)
    finally:
        workbook.close()


def _read_sheet(sheet) -> list[Direction]:
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_headers = list(next(rows))
    except StopIteration as exc:
        raise InputWorkbookError("Input workbook is empty") from exc

    headers = [_normalize_header(value) for value in raw_headers]
    point_col = _find_column(headers, SHIPPING_POINT_ALIASES, "Пункт отгрузки")
    region_col = _find_column(headers, DELIVERY_REGION_ALIASES, "Регион доставки")

    directions: list[Direction] = []
    seen: set[tuple[str, str]] = set()
    for excel_row, values_tuple in enumerate(rows, start=2):
        values = list(values_tuple)
        point = str(values[point_col] or "").strip()
        region = str(values[region_col] or "").strip()
        if not point and not region:
            continue
        if not point or not region:
            raise InputWorkbookError(
                f"Row {excel_row}: both shipping point and delivery region are required"
            )
        key = (point.casefold(), region.casefold())
        if key in seen:
            continue
        seen.add(key)
        extras = {
            str(raw_headers[i]).strip(): value
            for i, value in enumerate(values)
            if i not in {point_col, region_col}
            and i < len(raw_headers)
            and raw_headers[i] not in (None, "")
            and value not in (None, "")
        }
        directions.append(
            Direction(
                direction_id=f"DIR-{len(directions) + 1:04d}",
                shipping_point=point,
                delivery_region=region,
                attributes=extras,
            )
        )

    if not directions:
        raise InputWorkbookError("No valid directions found")
    return directions


def write_forecast(path: str | Path, response: ForecastResponse) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Прогноз"
    base_headers = [
        "ID направления",
        "Пункт отгрузки",
        "Регион доставки",
        "Период",
        "Прогноз",
    ]
    attribute_headers = list(
        dict.fromkeys(key for row in response.forecasts for key in row.attributes)
    )
    headers = base_headers + attribute_headers
    sheet.append(headers)
    for row in response.forecasts:
        sheet.append(
            [
                row.direction_id,
                row.shipping_point,
                row.delivery_region,
                row.period,
                row.forecast,
                *(
                    _excel_value(row.attributes.get(header))
                    for header in attribute_headers
                ),
            ]
        )

    # External text must remain text, never executable spreadsheet formulas.
    for cells in sheet.iter_rows():
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {"A": 18, "B": 28, "C": 30, "D": 14, "E": 14}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet["D"][1:]:
        cell.number_format = "mm.yyyy"
    for cell in sheet["E"][1:]:
        cell.number_format = "0.00"
    if sheet.max_column > 5:
        for column_cells in sheet.iter_cols(min_col=6, max_col=sheet.max_column):
            sheet.column_dimensions[column_cells[0].column_letter].width = 24
    if sheet.max_row > 1:
        table = Table(displayName="ForecastTable", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
        )
        sheet.add_table(table)
    workbook.save(output)
    workbook.close()
    return output


def _excel_value(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value
