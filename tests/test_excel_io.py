from datetime import UTC, date, datetime
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from shipment_forecast_agent.excel_io import (
    InputWorkbookError,
    read_directions,
    write_forecast,
)
from shipment_forecast_agent.models import ForecastResponse, ForecastRow


def test_read_directions_accepts_russian_headers_and_deduplicates(tmp_path):
    path = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Пункт отгрузки", "Регион доставки", "Комментарий"])
    sheet.append(["Малино", "Смоленская область", "A"])
    sheet.append(["малино", "смоленская область", "duplicate"])
    workbook.save(path)

    directions = read_directions(path)

    assert len(directions) == 1
    assert directions[0].shipping_point == "Малино"
    assert directions[0].attributes == {"Комментарий": "A"}


def test_read_directions_rejects_incomplete_row(tmp_path):
    path = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Пункт отгрузки", "Регион доставки"])
    sheet.append(["Малино", None])
    workbook.save(path)

    try:
        read_directions(path)
    except InputWorkbookError as exc:
        assert "Row 2" in str(exc)
    else:
        raise AssertionError("Expected InputWorkbookError")


def test_write_forecast(tmp_path):
    request_id = uuid4()
    response = ForecastResponse(
        request_id=request_id,
        generated_at=datetime.now(UTC),
        forecasts=[
            ForecastRow(
                direction_id="DIR-0001",
                shipping_point="Малино",
                delivery_region="Смоленская область",
                period=date(2026, 9, 1),
                forecast=9,
                attributes={"Код региона доставки": "67"},
            )
        ],
    )
    output = write_forecast(tmp_path / "result.xlsx", response)
    workbook = load_workbook(output, data_only=True)
    sheet = workbook["Прогноз"]
    assert sheet["B2"].value == "Малино"
    assert sheet["E2"].value == 9
    assert sheet["F1"].value == "Код региона доставки"
    assert sheet["F2"].value == "67"
    workbook.close()


def test_external_formula_stays_text(tmp_path):
    response = ForecastResponse(
        request_id=uuid4(),
        generated_at=datetime.now(UTC),
        forecasts=[
            ForecastRow(
                direction_id="D1",
                shipping_point="=1+1",
                delivery_region="Region",
                period=date(2026, 10, 1),
                forecast=1,
                attributes={"Details": {"a": 1}},
            )
        ],
    )
    output = write_forecast(tmp_path / "safe.xlsx", response)
    book = load_workbook(output)
    assert book.active["B2"].data_type == "s"
    assert book.active["B2"].value == "=1+1"
    assert book.active["F2"].value == '{"a": 1}'
    book.close()
