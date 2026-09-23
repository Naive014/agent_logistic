import pytest
from openpyxl import Workbook

from shipment_forecast_agent import agents
from shipment_forecast_agent.config import Settings
from shipment_forecast_agent.internet_mail import ImapInputEnvelope
from shipment_forecast_agent.template_excel import REQUIRED_TENDER_HEADERS, request_data


def save_template(path, headers):
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for index, label in enumerate(headers, 1):
        if label.strip().casefold() == "наименование пункта отгрузки":
            sheet.cell(2, index, "Завод")
        if label.strip().casefold() == "наименование региона доставки":
            sheet.cell(2, index, "Область")
    book.save(path)
    book.close()


@pytest.mark.parametrize("missing", REQUIRED_TENDER_HEADERS)
def test_every_template_column_is_required(tmp_path, missing):
    source = tmp_path / "input.xlsx"
    save_template(source, [h for h in REQUIRED_TENDER_HEADERS if h != missing])
    with pytest.raises(ValueError) as error:
        request_data(source, Settings(_env_file=None), "test")
    assert missing in str(error.value)
    assert "Отсутствуют" in str(error.value)


def test_order_case_whitespace_extra_columns_and_empty_optional_values(tmp_path):
    source = tmp_path / "input.xlsx"
    headers = ["  " + h.upper() + "\n" for h in reversed(REQUIRED_TENDER_HEADERS)]
    save_template(source, headers + ["Дополнительный столбец"])
    payload = request_data(source, Settings(_env_file=None), "test")
    assert payload["shipment_point_name"] == "Завод"
    assert payload["delivery_point_name"] == "Область"


@pytest.mark.parametrize("duplicate", ["Валюта", "Наименование пункта отгрузки"])
def test_duplicate_required_header_rejected(tmp_path, duplicate):
    source = tmp_path / "input.xlsx"
    save_template(source, [*REQUIRED_TENDER_HEADERS, duplicate.upper()])
    with pytest.raises(ValueError, match="Дублируются"):
        request_data(source, Settings(_env_file=None), "test")


def test_invalid_template_not_sent_or_marked_read(monkeypatch, tmp_path, caplog):
    source = tmp_path / "input.xlsx"
    save_template(source, list(REQUIRED_TENDER_HEADERS[:-1]))
    item = ImapInputEnvelope(
        b"1",
        "user@example.com",
        "[FORECAST_INPUT]",
        "input.xlsx",
        source.read_bytes(),
        "123",
    )
    monkeypatch.setattr(agents, "iter_input_workbooks", lambda s: [item])

    def unexpected(*args):
        pytest.fail("Invalid workbook must not be sent or marked read")

    monkeypatch.setattr(agents, "_send_now", unexpected)
    monkeypatch.setattr(agents, "mark_processed", unexpected)
    settings = Settings(_env_file=None, WORK_DIR=tmp_path / "state")
    assert agents.process_requests(settings) == []
    assert "Гарантированные объемы" in caplog.text
