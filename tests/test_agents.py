import json
from email.message import EmailMessage

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from shipment_forecast_agent import agents
from shipment_forecast_agent.config import Settings
from shipment_forecast_agent.internet_mail import ImapInputEnvelope
from shipment_forecast_agent.registry import Registry
from shipment_forecast_agent.template_excel import (
    REQUIRED_TENDER_HEADERS,
    current_month,
    fill_forecasts,
    request_data,
)


def workbook(path, periods=False, full_template=False):
    book = Workbook()
    sheet = book.active
    sheet.append(
        ["Пункт отгрузки", "Регион доставки", "Цена", "Расчёт"]
        + (["Период"] if periods else [])
    )
    sheet.append(["Завод", "Область", 100, "=C2*2"] + (["2026-09"] if periods else []))
    if periods:
        sheet.append(["Завод", "Область", 200, "=C3*2", "2026-10"])
    if full_template:
        for label in REQUIRED_TENDER_HEADERS:
            if label not in (
                "Наименование пункта отгрузки",
                "Наименование региона доставки",
            ):
                sheet.cell(1, sheet.max_column + 1, label)
    sheet["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    book.create_sheet("Справочник")["A1"] = "Сохранить"
    book.save(path)
    book.close()


def forecast(period="2026-09", value=10):
    return {
        "shipment_point_name": "Завод",
        "delivery_point_name": "Область",
        "ppr_month": period,
        "forecast_m1": value,
        "forecast_m2": value + 1,
        "forecast_m3": value + 2,
    }


def test_split_agents_roundtrip_preserves_source(monkeypatch, tmp_path):
    source = tmp_path / "request.xlsx"
    workbook(source, full_template=True)
    original = source.read_bytes()
    settings = Settings(
        _env_file=None,
        WORK_DIR=tmp_path / "state",
        MAILBOX_NAME="bridge@example.com",
        INPUT_MAILBOX_NAME="planning@example.com",
        INPUT_MAILBOX_PASSWORD="secret",
        INTERNAL_OUTLOOK_EMAIL="corp@example.com",
        FORECAST_START_MONTH="2026-09",
    )
    item = ImapInputEnvelope(
        b"42",
        "author@example.com",
        "[FORECAST_INPUT]",
        "safe.xlsx",
        original,
        "123",
        "ТЗ.xlsx",
        "2026-09-22T10:00:00+00:00",
        "<original@example.com>",
        "<older@example.com>",
    )
    monkeypatch.setattr(agents, "iter_input_workbooks", lambda s: [item])
    sent, marked = [], []
    monkeypatch.setattr(agents, "_send_now", lambda s, m: sent.append(m))
    monkeypatch.setattr(
        agents, "mark_processed", lambda s, uid: marked.append((s.MAILBOX_NAME, uid))
    )
    (key,) = agents.process_requests(settings)
    assert sent == []
    record = Registry(settings.WORK_DIR).get(key)
    assert record["status"] == "prepared"
    assert record["outlook_sent_at"] == ""
    assert record["filename"] == "ТЗ.xlsx" and record["mail_uid"] == "42"
    assert record["original_subject"] == "[FORECAST_INPUT]"
    assert record["original_message_id"] == "<original@example.com>"

    response_messages = []

    def responses(_, prefix, *, unseen_only):
        assert prefix is None
        assert unseen_only is False
        return list(response_messages)

    monkeypatch.setattr(agents, "_iter_messages", responses)
    assert agents.process_results(settings) == [key]
    assert sent[0]["Subject"] == "[N8N_FORECAST_REQUEST] " + key
    payload = json.loads(next(sent[0].iter_attachments()).get_payload(decode=True))
    assert payload == {
        "mail_id": key,
        "param_month_m0": current_month(),
        "shipment_point_name": "Завод",
        "delivery_point_name": "Область",
    }
    assert Registry(settings.WORK_DIR).get(key)["status"] == "sent"
    assert agents.process_requests(settings) == []
    assert len(sent) == 1
    msg = EmailMessage()
    msg["From"] = "corp@example.com"
    msg["X-Agent-Received-At"] = "2026-09-22T10:05:00+00:00"
    msg["Subject"] = "n8n result without correlation ID"
    msg.set_content("Result")
    msg.add_attachment(
        json.dumps(
            dict(forecast(), mail_id=key, reply_to="untrusted@example.com")
        ).encode(),
        maintype="application",
        subtype="json",
        filename="result.json",
    )
    response_messages.append((b"99", "555", msg))
    results = []
    monkeypatch.setattr(
        agents,
        "send_result_workbook",
        lambda s, recipient, uid, path, **kwargs: results.append(
            (s.MAILBOX_NAME, recipient, uid, path.name, kwargs)
        ),
    )
    (output,) = agents.process_results(settings)
    assert results == [
        (
            "planning@example.com",
            "author@example.com",
            key,
            "ТЗ.xlsx",
            {
                "original_subject": "[FORECAST_INPUT]",
                "original_message_id": "<original@example.com>",
                "original_references": "<older@example.com>",
                "attachment_filename": "ТЗ.xlsx",
            },
        )
    ]
    book = load_workbook(output)
    assert [book.active.cell(1, column).value for column in (24, 25, 26)] == [
        "Прогноз M1",
        "Прогноз M2",
        "Прогноз M3",
    ]
    assert output.name == "ТЗ.xlsx"
    assert book.active.max_column == 26
    assert [book.active.cell(2, column).value for column in (24, 25, 26)] == [
        10,
        11,
        12,
    ]
    assert book.active["D2"].value == "=C2*2"
    assert book.active["A1"].fill.fgColor.rgb == "00FFFF00"
    assert book["Справочник"]["A1"].value == "Сохранить"
    book.close()
    assert source.read_bytes() == original
    assert (settings.WORK_DIR / "jobs" / key / "source.xlsx").read_bytes() == original
    assert Registry(settings.WORK_DIR).get(key)["status"] == "completed"
    audit = Registry(settings.WORK_DIR).get(key)
    assert audit["UID"] == key
    assert audit["request_received_at"] == "2026-09-22T10:00:00+00:00"
    assert audit["outlook_received_at"] == "2026-09-22T10:05:00+00:00"
    assert audit["outlook_sent_at"] and audit["user_sent_at"]
    assert audit["output_filename"] == "ТЗ.xlsx"
    assert agents.process_results(settings) == []
    assert len(results) == 1


def test_same_direction_gets_one_value_regardless_of_period(tmp_path):
    source, output = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    workbook(source, periods=True)
    fill_forecasts(source, output, [forecast()], Settings(_env_file=None))
    book = load_workbook(output)
    assert book.active["F2"].value == 10
    assert book.active["F3"].value == 10
    assert book.active["G2"].value == 11
    assert book.active["G3"].value == 11
    assert book.active["H2"].value == 12
    assert book.active["H3"].value == 12
    book.close()


def test_ambiguous_months_do_not_produce_output(tmp_path):
    source, output = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    workbook(source)
    with pytest.raises(ValueError, match="ambiguous"):
        fill_forecasts(
            source, output, [forecast(), forecast("2026-10")], Settings(_env_file=None)
        )
    assert not output.exists()


def test_missing_or_conflicting_uid_rejected():
    key = "RQ-" + "a" * 32
    with pytest.raises(ValueError):
        agents.correlate("[N8N_FORECAST_RESULT]", forecast())
    with pytest.raises(ValueError):
        agents.correlate(key, {"mail_id": "RQ-" + "b" * 32})
    with pytest.raises(ValueError):
        agents.correlate(key, forecast())
    assert agents.correlate(key, dict(forecast(), mail_id=key)) == key
    assert agents.correlate("Result", {"mail_id": key}) == key


def test_mail_id_in_rows_and_conflicts():
    key = "RQ-" + "a" * 32
    assert agents.correlate("Result", {"data": [dict(forecast(), mail_id=key)]}) == key
    with pytest.raises(ValueError):
        agents.correlate(
            "Result", {"mail_id": key, "data": [{"mail_id": "RQ-" + "b" * 32}]}
        )


def test_empty_response_rejected():
    with pytest.raises(ValueError):
        agents.forecast_rows({"status": "ok", "row_count": 0, "data": []})


def test_terminal_invalid_response_is_only_checked_once(monkeypatch, tmp_path):
    key = "RQ-" + "a" * 32
    settings = Settings(
        _env_file=None,
        WORK_DIR=tmp_path,
        MAILBOX_NAME="bridge@example.com",
        INTERNAL_OUTLOOK_EMAIL="corp@example.com",
    )
    message = EmailMessage()
    message["From"] = "corp@example.com"
    message["Subject"] = key
    message.set_content("Empty result")
    message.add_attachment(
        json.dumps(
            {"mail_id": key, "status": "ok", "row_count": 0, "data": []}
        ).encode(),
        maintype="application",
        subtype="json",
        filename="result.json",
    )
    calls = []

    def responses(*args, **kwargs):
        calls.append((args, kwargs))
        return [(b"9", "555", message)]

    monkeypatch.setattr(agents, "_iter_messages", responses)
    assert agents.process_results(settings) == []
    assert agents.process_results(settings) == []
    handled = json.loads(
        (tmp_path / "result-mail-handled.json").read_text(encoding="utf-8")
    )
    assert handled == {"handled": ["555:9"]}
    assert len(calls) == 2


@pytest.mark.parametrize(
    "year,month,expected", [(2026, 9, "2026-09"), (2027, 1, "2027-01")]
)
def test_request_uses_current_calendar_month(
    monkeypatch, tmp_path, year, month, expected
):
    from datetime import date

    from shipment_forecast_agent import template_excel

    class Clock:
        @staticmethod
        def today():
            return date(year, month, 1)

    monkeypatch.setattr(template_excel, "date", Clock)
    source = tmp_path / "source.xlsx"
    workbook(source, periods=True, full_template=True)
    payload = template_excel.request_data(
        source,
        Settings(_env_file=None, FORECAST_START_MONTH="2000-01"),
        "RQ-" + "a" * 32,
    )
    assert payload["param_month_m0"] == expected
    assert set(payload) == {
        "mail_id",
        "param_month_m0",
        "shipment_point_name",
        "delivery_point_name",
    }


def test_request_rejects_multiple_unique_directions(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook(source, full_template=True)
    book = load_workbook(source)
    book.active.append(["Другой завод", "Другая область"])
    book.save(source)
    book.close()
    with pytest.raises(ValueError, match="exactly one unique direction"):
        request_data(source, Settings(_env_file=None), "RQ-" + "a" * 32)


def test_forecast_without_month(tmp_path):
    source, output = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    workbook(source)
    result = forecast()
    del result["ppr_month"]
    fill_forecasts(source, output, [result], Settings(_env_file=None))
    book = load_workbook(output)
    assert book.active["E2"].value == 10
    assert book.active["F2"].value == 11
    assert book.active["G2"].value == 12
    book.close()


def test_new_n8n_response_uses_all_three_forecasts(tmp_path):
    source, output = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    workbook(source)
    result = {
        "shipment_point_name": "Завод",
        "delivery_point_name": "Область",
        "param_month_m0": "2026-09",
        "forecast_m1": 31.49580146836055,
        "forecast_m2": 45.6,
        "forecast_m3": 45.6,
    }
    fill_forecasts(source, output, [result], Settings(_env_file=None))
    book = load_workbook(output)
    assert book.active["E2"].value == pytest.approx(31.49580146836055)
    assert book.active["F2"].value == pytest.approx(45.6)
    assert book.active["G2"].value == pytest.approx(45.6)
    book.close()


@pytest.mark.parametrize("missing", ["forecast_m1", "forecast_m2", "forecast_m3"])
def test_missing_forecast_period_is_rejected(tmp_path, missing):
    source, output = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    workbook(source)
    result = forecast()
    del result[missing]

    with pytest.raises(ValueError, match=missing):
        fill_forecasts(source, output, [result], Settings(_env_file=None))

    assert not output.exists()


def test_registry_update_keeps_other_requests(tmp_path):
    registry = Registry(tmp_path)
    registry.put({"request_id": "a", "filename": "А.xlsx", "status": "prepared"})
    registry.put({"request_id": "b", "filename": "Б.xlsx", "status": "sent"})
    registry.put({"request_id": "a", "status": "completed"})
    assert registry.get("a")["filename"] == "А.xlsx"
    assert registry.get("b")["status"] == "sent"


def test_result_agent_sends_only_one_prepared_request_per_cycle(monkeypatch, tmp_path):
    settings = Settings(
        _env_file=None,
        WORK_DIR=tmp_path,
        MAILBOX_NAME="bridge@example.com",
        INTERNAL_OUTLOOK_EMAIL="corp@example.com",
    )
    registry = Registry(tmp_path)
    keys = ["RQ-" + "a" * 32, "RQ-" + "b" * 32]
    for key in keys:
        directory = tmp_path / "jobs" / key
        directory.mkdir(parents=True)
        (directory / "request.json").write_text(
            json.dumps(
                {
                    "mail_id": key,
                    "param_month_m0": "2026-09",
                    "shipment_point_name": "Завод",
                    "delivery_point_name": "Область",
                }
            ),
            encoding="utf-8",
        )
        registry.put({"request_id": key, "status": "prepared"})

    sent = []
    monkeypatch.setattr(agents, "_iter_messages", lambda *a, **kw: [])
    monkeypatch.setattr(agents, "_send_now", lambda s, m: sent.append(m))

    assert agents.process_results(settings) == [keys[0]]
    assert len(sent) == 1
    assert registry.get(keys[0])["status"] == "sent"
    assert registry.get(keys[1])["status"] == "prepared"

    assert agents.process_results(settings) == [keys[1]]
    assert len(sent) == 2
    assert registry.get(keys[1])["status"] == "sent"


def test_request_agent_takes_only_one_source_email_per_cycle(monkeypatch, tmp_path):
    source = tmp_path / "request.xlsx"
    workbook(source, full_template=True)
    items = [
        ImapInputEnvelope(
            str(uid).encode(),
            "author@example.com",
            "[FORECAST_INPUT]",
            f"safe-{uid}.xlsx",
            source.read_bytes(),
            "123",
            f"request-{uid}.xlsx",
        )
        for uid in (1, 2)
    ]
    marked = []
    monkeypatch.setattr(agents, "iter_input_workbooks", lambda settings: items)
    monkeypatch.setattr(
        agents, "mark_processed", lambda settings, uid: marked.append(uid)
    )
    settings = Settings(_env_file=None, WORK_DIR=tmp_path / "state")

    prepared = agents.process_requests(settings)

    assert len(prepared) == 1
    assert marked == [b"1"]
    assert len(list((settings.WORK_DIR / "jobs").iterdir())) == 1
